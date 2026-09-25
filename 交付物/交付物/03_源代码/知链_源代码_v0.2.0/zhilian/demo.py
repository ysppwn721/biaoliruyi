"""Synthetic, explicitly labelled demo; the application modifies real Office files."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from .office import HEADERS

ROWS = [
    ['sales_prev', '总计', '销售额', '上期', 100, '万元', '演示业务'],
    ['sales_current', '总计', '销售额', '本期', 125, '万元', '演示业务'],
    ['product_a', 'A产品', '销量', '本期', 80, '件', '演示产品'],
    ['product_b', 'B产品', '销量', '本期', 65, '件', '演示产品'],
    ['spending', '总计', '支出', '本期', 80, '万元', '演示项目'],
    ['budget', '总计', '预算', '本期', 100, '万元', '演示项目'],
]


def create_demo(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    sheet = wb.active
    sheet.title = '事实表'
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in zip('ABCDEFG', [24, 15, 15, 14, 14, 12, 22]):
        sheet.column_dimensions[col].width = width
    for cell in sheet[1]:
        cell.font = Font(name='Microsoft YaHei', color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='155D50')
    for row in sheet.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name='Microsoft YaHei', size=11)
            c.alignment = Alignment(vertical='center')
        sheet.row_dimensions[row[0].row].height = 25
    wb.save(folder / '业务数据.xlsx')
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    from docx.shared import RGBColor as WordRGB
    for name in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[name].font.color.rgb = WordRGB(0, 0, 0)
        doc.styles[name].font.name = 'Microsoft YaHei'
        rpr = doc.styles[name].element.get_or_add_rPr()
        rpr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.add_heading('业务分析报告', 0)
    doc.add_paragraph('模拟数据 · 用于验证知链功能，不代表真实经营结果。')
    doc.add_heading('销售表现', 1)
    para = doc.add_paragraph()
    para.add_run('本期销售额为').bold = True
    para.add_run('125').italic = True
    para.add_run('万元。较上期增长25%。本期销售额超过120万元。')
    doc.add_heading('产品与预算', 1)
    doc.add_paragraph('A产品销量最高。支出未超过预算。')
    doc.add_paragraph('本段是人工撰写的背景说明，修复时应保持原样。')
    doc.save(folder / '分析报告.docx')
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def slide(title, lines):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string('F4F7F5')
        tf = s.shapes.add_textbox(Inches(.7), Inches(.6), Inches(11.9), Inches(.7)).text_frame
        tf.text = title
        for r in tf.paragraphs[0].runs:
            r.font.size = PptPt(30); r.font.bold = True; r.font.name = 'Microsoft YaHei'
        if lines:
            tf = s.shapes.add_textbox(Inches(.8), Inches(1.8), Inches(11.7), Inches(4.6)).text_frame
            for i, line in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                p.space_after = PptPt(26)
                for r in p.runs:
                    r.font.size = PptPt(24); r.font.name = 'Microsoft YaHei'
        return s
    slide('业务汇报', ['模拟数据 · 结论来自业务数据.xlsx', '本期销售额为125万元。较上期增长25%。', '本期销售额超过120万元。'])
    slide('产品与预算', ['A产品销量最高。', '支出未超过预算。', '此处为人工补充的说明，保持不变。'])
    s = slide('销售额对比  单位万元', [])
    data = CategoryChartData()
    data.categories = ['上期', '本期']
    data.add_series('销售额（万元）', [100, 125])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.7), Inches(11), Inches(4.7), data).chart
    chart.has_legend = False
    chart.value_axis.minimum_scale = 0
    prs.save(folder / '业务汇报.pptx')
    return [folder / '业务数据.xlsx', folder / '分析报告.docx', folder / '业务汇报.pptx']
