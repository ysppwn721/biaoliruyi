from copy import deepcopy

import pytest
from openpyxl import load_workbook

from zhilian.demo import create_demo
from zhilian.engine import check
from zhilian.store import Store


def confirmed_project(tmp_path):
    store = Store(tmp_path / 'data')
    store.create('预览项目', create_demo(tmp_path / 'files'), True)
    state = store.read(store.list()[0]['id'])
    result = store.confirm(state['id'], state['revision'],
                           [{'claim_id': c['id'], 'refs': c['refs']} for c in state['claims']])
    return store, result


def updated_book(tmp_path, value=123):
    path = create_demo(tmp_path / 'external')[0]
    wb = load_workbook(path)
    wb.active['E3'] = value
    wb.save(path)
    wb.close()
    return path


def test_preview_is_read_only_and_keeps_true_claims(tmp_path):
    store, ws = confirmed_project(tmp_path)
    before = store.read(ws['id'])
    generation = before['generation']
    files = {p.name: p.stat().st_mtime_ns for p in (store.folder(ws['id']) / generation).iterdir()}

    result = store.preview_source(ws['id'], ws['revision'], updated_book(tmp_path, 123))

    assert result['revision'] == ws['revision']
    assert result['summary']['changed_facts'] == 1
    assert result['summary']['affected_claims'] > 0
    assert any(a['after_status'] == 'inconsistent' for a in result['affected'])
    assert all(a['claim_id'] in {c['id'] for c in ws['claims']} for a in result['affected'])
    assert store.read(ws['id']) == before
    assert {p.name: p.stat().st_mtime_ns for p in (store.folder(ws['id']) / generation).iterdir()} == files


def test_preview_rejects_conflict_and_semantic_change(tmp_path):
    store, ws = confirmed_project(tmp_path)
    path = updated_book(tmp_path, 123)
    with pytest.raises(ValueError, match='其他窗口'):
        store.preview_source(ws['id'], ws['revision'] - 1, path)

    wb = load_workbook(path)
    wb.active['G3'] = '其他口径'
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match='统计口径'):
        store.preview_source(ws['id'], ws['revision'], path)


def test_preview_matches_applied_import_checks(tmp_path):
    store, ws = confirmed_project(tmp_path)
    path = updated_book(tmp_path, 123)
    preview = store.preview_source(ws['id'], ws['revision'], path)
    updated = store.import_source(ws['id'], ws['revision'], path)
    actual = {c['id']: check(c, updated['facts']) for c in updated['claims']}
    assert {a['claim_id']: a['after_status'] for a in preview['affected']} == {
        cid: actual[cid]['status'] for cid in (a['claim_id'] for a in preview['affected'])
    }
