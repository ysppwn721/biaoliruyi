import base64
import json
from copy import deepcopy

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from zhilian import app as app_module, ocr
from zhilian.agent import run_agent
from zhilian.demo import create_demo
from zhilian.office import digest, read_images, read_document
from zhilian.store import Store


@pytest.fixture(autouse=True)
def isolated(monkeypatch):
    monkeypatch.setattr(app_module, 'load_environment', lambda: None)
    monkeypatch.setenv('DEEPSEEK_API_KEY', '')
    monkeypatch.setenv('DEEPSEEK_MODEL', 'test-vision')
    monkeypatch.setenv('ZHILIAN_OCR_BACKEND', 'auto')
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', '')


@pytest.fixture
def materials(tmp_path):
    paths = create_demo(tmp_path / 'files')
    picture = tmp_path / 'sample.png'
    image = Image.new('RGB', (480, 100), 'white')
    ImageDraw.Draw(image).text((12, 20), 'Sales 125', fill='black')
    image.save(picture)
    doc = Document(paths[1])
    doc.add_picture(str(picture))
    doc.save(paths[1])
    deck = Presentation(paths[2])
    deck.slides[0].shapes.add_picture(str(picture), Inches(1), Inches(1))
    deck.save(paths[2])
    return paths, picture


@pytest.fixture
def project(tmp_path, materials):
    store = Store(tmp_path / 'data')
    return store, store.create('OCR 测试', materials[0])


@pytest.fixture
def api_project(tmp_path, materials):
    app = app_module.create_app(tmp_path / 'api')
    with TestClient(app) as client:
        response = client.post('/api/projects', data={'name': '含图项目'},
                               files=[('files', (p.name, p.read_bytes())) for p in materials[0]])
        assert response.status_code == 200
        yield client, app.state.store, response.json()


def response_text(text='本期销售额为125万元。'):
    return httpx.Response(200, json={'choices': [{'message': {'content': text}, 'finish_reason': 'stop'}]})


def hashes(store, ws):
    return {d['stored_name']: digest(store.folder(ws['id']) / ws['generation'] / d['stored_name'])
            for d in ws['documents'] + ws.get('images', [])}


def enable(monkeypatch):
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-secret')
    monkeypatch.setattr(httpx, 'post', lambda *a, **k: response_text())


def test_extract_images_and_preserve_reader(project, materials):
    store, ws = project
    assert len(ws['images']) == 2
    assert ws['ocr_blocks'] == ws['ocr_claims'] == []
    stored = store.read(ws['id'])
    assert {i['file_id'] for i in ws['images']} == {d['id'] for d in ws['documents'] if d['kind'] != 'xlsx'}
    assert {json.loads(i['location'])[0] for i in ws['images']} == {'inline', 'image'}
    for item in ws['images']:
        assert item['id'].startswith('img') and len(item['id']) == 19
        assert item['mime'] == 'image/png' and item['label']
        assert item['ocr_status'] == 'pending' and item['ocr_text'] is None
        assert 'stored_name' not in item and 'blob' not in item
        saved = next(i for i in stored['images'] if i['id'] == item['id'])
        assert store.image_path(stored, saved).read_bytes() == materials[1].read_bytes()
        assert item['sha256'] == digest(materials[1])
    for path in materials[0][1:]:
        before = read_document(path, 'file', ws['facts'])
        read_images(path, 'file')
        assert read_document(path, 'file', ws['facts']) == before


def test_image_only_import_and_group_warning(tmp_path, materials):
    doc = Document()
    doc.add_picture(str(materials[1]))
    doc.save(materials[0][1])
    store = Store(tmp_path / 'only')
    ws = store.create('仅图片', materials[0][:2])
    assert ws['images'] and ws['blocks'] == [] and ws['claims'] == []
    deck = Presentation(materials[0][2])
    group = deck.slides[0].shapes.add_group_shape()
    group.shapes.add_picture(str(materials[1]), Inches(1), Inches(1))
    deck.save(materials[0][2])
    images, warnings = read_images(materials[0][2], 'ppt')
    assert len(images) == 1
    assert any('组合对象' in w for w in warnings)


def test_multimodal_payload_and_gate(api_project, materials, monkeypatch):
    client, store, ws = api_project
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-secret')
    before = store.read(ws['id'])
    old_hashes = hashes(store, before)
    calls = []
    def post(url, **kwargs):
        calls.append(url)
        assert url == 'https://api.deepseek.com/chat/completions'
        payload = kwargs['json']
        assert payload['model'] == 'test-vision'
        assert payload['max_tokens'] == 4000 and payload['temperature'] == 0
        assert payload['stream'] is False and kwargs['follow_redirects'] is False
        assert 'response_format' not in payload and 'thinking' not in payload
        content = payload['messages'][0]['content']
        uri = content[0]['image_url']['url']
        assert uri.startswith('data:image/png;base64,')
        assert base64.b64decode(uri.split(',', 1)[1]) == materials[1].read_bytes()
        assert '逐字识别' in content[1]['text']
        return response_text()
    monkeypatch.setattr(httpx, 'post', post)
    endpoint = f'/api/projects/{ws["id"]}/ocr'
    pending = client.post(endpoint + '/confirm', json={'revision': ws['revision'], 'items': [
        {'image_id': ws['images'][0]['id'], 'text': '本期销售额为125万元。'}]})
    assert pending.status_code == 400
    response = client.post(endpoint + '/run', json={'revision': ws['revision']})
    assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
    done = response.json()
    assert len(calls) == 2 and all(i['ocr_status'] == 'done' for i in done['images'])
    assert done['ocr_claims'] == done['ocr_blocks'] == []
    assert done['claims'] == ws['claims'] and done['checks'] == ws['checks']
    assert hashes(store, store.read(ws['id'])) == old_hashes
    assert 'test-secret' not in response.text
    response = client.post(endpoint + '/confirm', json={'revision': done['revision'], 'items': [
        {'image_id': done['images'][0]['id'], 'text': '本期销售额为100万元。'}]})
    confirmed = response.json()
    assert response.status_code == 200
    claim = confirmed['ocr_claims'][0]
    assert claim['source'] == 'ocr' and claim['repairable'] is False and claim['confirmed'] is False
    assert claim['original'] == '本期销售额为100万元'
    checked = next(r for r in confirmed['checks'] if r['claim_id'] == claim['id'])
    assert checked['ocr'] is True and checked['status'] == 'inconsistent'
    diagnosis = next(r for r in confirmed['diagnosis'] if r['claim_id'] == claim['id'])
    assert diagnosis['file_name'] and diagnosis['source'] == 'ocr' and diagnosis['evidence'][0]['cell'] == 'E3'
    assert '图片论断仅用于只读验证' in diagnosis['fix_hint']
    assert confirmed['summary']['claims'] == ws['summary']['claims'] + 1
    assert confirmed['summary']['repairable'] == ws['summary']['repairable']
    assert confirmed['summary']['pending'] == ws['summary']['pending'] + 1
    assert store.read(ws['id'])['claims'] == before['claims']
    assert hashes(store, store.read(ws['id'])) == old_hashes
    assert 'test-secret' not in json.dumps(store.read(ws['id']))


@pytest.mark.parametrize('backend,key,enabled,resolved', [
    ('auto', '', False, 'off'), ('auto', 'test-secret', True, 'deepseek'),
    ('off', 'test-secret', False, 'off'), ('deepseek', '', False, 'deepseek'),
    ('deepseek', 'test-secret', True, 'deepseek'), ('invalid', 'test-secret', False, 'off')])
def test_config(backend, key, enabled, resolved, monkeypatch):
    monkeypatch.setenv('ZHILIAN_OCR_BACKEND', backend)
    monkeypatch.setenv('DEEPSEEK_API_KEY', key)
    assert ocr.config() == {'enabled': enabled, 'backend': resolved, 'model': 'test-vision', 'key_configured': bool(key)}


def test_disabled_no_network_or_mutation(api_project, monkeypatch):
    client, store, ws = api_project
    monkeypatch.setenv('ZHILIAN_OCR_BACKEND', 'off')
    def forbidden(*a, **k):
        pytest.fail('OCR 关闭时不得发请求')
    monkeypatch.setattr(httpx, 'post', forbidden)
    before = store.read(ws['id'])
    result = client.post(f'/api/projects/{ws["id"]}/ocr/run', json={'revision': ws['revision']})
    assert result.status_code == 200 and result.json()['images'][0]['ocr_status'] == 'pending'
    assert store.read(ws['id']) == before
    assert client.get('/api/health').json()['ocr']['enabled'] is False
    with pytest.raises(ValueError, match='未启用'):
        ocr.ocr_image('unused', 'image/png')


@pytest.mark.parametrize('error', ['401', '500', 'timeout', 'json', 'empty', 'leak', 'truncated'])
def test_failure_safe_and_retry(api_project, monkeypatch, error):
    client, store, ws = api_project
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-secret')
    def post(*a, **k):
        if error in ('401', '500'):
            return httpx.Response(int(error), text='provider-secret test-secret')
        if error == 'timeout':
            raise httpx.ReadTimeout('provider-secret test-secret')
        if error == 'json':
            return httpx.Response(200, text='provider-secret test-secret')
        if error == 'truncated':
            return httpx.Response(200, json={'choices': [{'message': {'content': '本期销售额为100万元'}, 'finish_reason': 'length'}]})
        return response_text('test-secret' if error == 'leak' else '')
    monkeypatch.setattr(httpx, 'post', post)
    before = store.read(ws['id'])
    response = client.post(f'/api/projects/{ws["id"]}/ocr/run', json={'revision': ws['revision']})
    assert response.status_code == 200
    result = response.json()
    assert all(i['ocr_status'] == 'failed' and i['ocr_error'] for i in result['images'])
    assert result['claims'] == ws['claims'] and result['facts'] == ws['facts'] and result['ocr_claims'] == []
    assert 'provider-secret' not in response.text and 'test-secret' not in response.text
    assert hashes(store, store.read(ws['id'])) == hashes(store, before)
    enable(monkeypatch)
    result = store.ocr_run(ws['id'], result['revision'])
    assert all(i['ocr_status'] == 'done' for i in result['images'])


def test_readonly_and_version_roundtrip(project, monkeypatch):
    store, ws = project
    enable(monkeypatch)
    ws = store.ocr_run(ws['id'], ws['revision'])
    iid = ws['images'][0]['id']
    ws = store.ocr_confirm(ws['id'], ws['revision'], [{'image_id': iid, 'text': '本期销售额为125万元。'}])
    cid = ws['ocr_claims'][0]['id']
    before = store.read(ws['id'])
    for operation in (lambda: store.confirm(ws['id'], ws['revision'], [{'claim_id': cid, 'refs': ['sales_current']}]),
                      lambda: store.repair(ws['id'], ws['revision'], [cid])):
        with pytest.raises(ValueError):
            operation()
        assert store.read(ws['id']) == before
    ws = store.confirm(ws['id'], ws['revision'], [{'claim_id': c['id'], 'refs': c['refs']} for c in before['claims']])
    ws = store.change(ws['id'], ws['revision'], {'sales_current': 90})
    assert next(r for r in ws['checks'] if r.get('ocr'))['status'] == 'inconsistent'
    original_images = deepcopy(ws['images'])
    ids = [r['claim_id'] for r in ws['checks'] if r['status'] == 'inconsistent' and not r.get('ocr')]
    ws = store.repair(ws['id'], ws['revision'], ids)
    assert ws['summary']['inconsistent'] == 1 and ws['summary']['repairable'] == 0
    assert ws['images'] == original_images
    ws = run_agent(store, ws['id'], ws['revision'])
    assert ws['agent']['phase'] == 'done' and ws['summary']['inconsistent'] == 1
    assert ws['ocr_claims'] == before['ocr_claims']
    ws = store.undo(ws['id'], ws['revision'])
    ws = store.undo(ws['id'], ws['revision'])
    assert next(r for r in ws['checks'] if r.get('ocr'))['status'] == 'consistent'
    assert store.image_path(store.read(ws['id']), store.read(ws['id'])['images'][0]).is_file()


def test_reconfirm_replaces_and_unverifiable_needs_review(project, monkeypatch):
    store, ws = project
    enable(monkeypatch)
    ws = store.ocr_run(ws['id'], ws['revision'])
    iid = ws['images'][0]['id']
    ws = store.ocr_confirm(ws['id'], ws['revision'], [{'image_id': iid, 'text': '本期销售额为125万元。'}])
    ws = store.ocr_confirm(ws['id'], ws['revision'], [{'image_id': iid, 'text': '未知指标为125万元。'}])
    assert len(ws['ocr_claims']) == len(ws['ocr_blocks']) == 1
    assert next(r for r in ws['checks'] if r.get('ocr'))['code'] == 'needs_review'
    assert next(r for r in ws['diagnosis'] if r.get('source') == 'ocr')['category'] == '需人工复核'
    ws = store.ocr_confirm(ws['id'], ws['revision'], [{'image_id': iid, 'text': '普通说明文字'}])
    assert ws['ocr_claims'] == []
    assert any(s.get('source') == 'ocr' for s in ws['unmatched_segments'])


def test_image_download_security_and_tamper(api_project, monkeypatch):
    client, store, ws = api_project
    image = ws['images'][0]
    response = client.get(image['image_url'])
    assert response.status_code == 200 and response.content
    assert response.headers['cache-control'] == 'no-store'
    assert response.headers['content-type'] == 'image/png'
    assert client.get(f'/api/projects/{ws["id"]}/images/missing').status_code == 404
    endpoint = f'/api/projects/{ws["id"]}/ocr/run'
    assert client.post(endpoint, json={'revision': ws['revision']}, headers={'origin': 'https://evil.test'}).status_code == 403
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', 'password')
    assert client.get(image['image_url']).status_code == 401
    assert client.post(endpoint, json={'revision': ws['revision']}).status_code == 401
    assert client.get(image['image_url'], auth=('zhilian', 'password')).status_code == 200
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', '')
    state = store.read(ws['id'])
    path = store.image_path(state, state['images'][0])
    path.write_bytes(b'tampered')
    assert client.get(image['image_url']).status_code == 400
    enable(monkeypatch)
    with pytest.raises(ValueError, match='图片'):
        store.ocr_run(ws['id'], ws['revision'])
    assert store.read(ws['id']) == state


def test_revision_guards_and_atomic_confirmation(project, monkeypatch):
    store, ws = project
    enable(monkeypatch)
    def racing(*a, **k):
        current = store.read(ws['id'])
        store.change(ws['id'], current['revision'], {'sales_current': 90 if current['revision'] == 0 else 80})
        return response_text()
    monkeypatch.setattr(httpx, 'post', racing)
    with pytest.raises(ValueError, match='其他窗口'):
        store.ocr_run(ws['id'], ws['revision'])
    state = store.read(ws['id'])
    assert all(i['ocr_status'] == 'pending' for i in state['images'])
    enable(monkeypatch)
    ws = store.ocr_run(ws['id'], state['revision'], [state['images'][0]['id']])
    before = store.read(ws['id'])
    with pytest.raises(ValueError, match='先运行'):
        store.ocr_confirm(ws['id'], ws['revision'], [{'image_id': i['id'], 'text': '本期销售额为100万元。'} for i in ws['images']])
    assert store.read(ws['id']) == before
    with pytest.raises(ValueError, match='其他窗口'):
        store.ocr_confirm(ws['id'], 0, [{'image_id': ws['images'][0]['id'], 'text': '修改'}])


def test_old_project_compatible(project):
    store, ws = project
    state = store.read(ws['id'])
    for key in ('images', 'ocr_blocks', 'ocr_claims'):
        state.pop(key)
    store.write(state)
    assert store.public(state)['images'] == []
    assert store.ocr_run(state['id'], state['revision'])['claims'] == state['claims']
