from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from pptx import Presentation

from zhilian.demo import create_demo
from zhilian.store import Store


def test_repair_chart_with_missing_embedded_workbook(tmp_path):
    paths = create_demo(tmp_path / 'files')
    deck = next(p for p in paths if p.suffix == '.pptx')
    # 模拟图表缓存仍可读、关系指向的嵌入工作簿已经丢失的真实文件。
    with ZipFile(deck) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()
                   if not info.filename.startswith('ppt/embeddings/')]
    with ZipFile(deck, 'w', ZIP_DEFLATED) as archive:
        for info, data in entries:
            archive.writestr(info, data)
    store = Store(tmp_path / 'data')
    ws = store.create('图表工作簿缺失测试', paths)
    ws = store.confirm(ws['id'], ws['revision'], [
        {'claim_id': c['id'], 'refs': c['refs']} for c in ws['claims']])
    ws = store.change(ws['id'], ws['revision'], {'sales_current': 90})
    raw = store.read(ws['id'])
    folder = store.folder(ws['id']) / raw['generation']
    original = {d['id']: (folder / d['stored_name']).read_bytes() for d in raw['documents']}
    ids = [c['claim_id'] for c in ws['checks'] if c['status'] == 'inconsistent']
    ws = store.repair(ws['id'], ws['revision'], ids)
    assert ws['last_repair']['verified'] and ws['summary']['inconsistent'] == 0
    raw = store.read(ws['id'])
    output = store.folder(ws['id']) / raw['generation']
    ppt = next(d for d in raw['documents'] if d['kind'] == 'pptx')
    prs = Presentation(output / ppt['stored_name'])
    chart = next(s.chart for slide in prs.slides for s in slide.shapes if s.has_chart)
    assert list(chart.series[0].values) == [100, 90]
    workbook = load_workbook(BytesIO(chart.part.chart_workbook.xlsx_part.blob))
    assert [workbook.active['B2'].value, workbook.active['B3'].value] == [100, 90]
    workbook.close()
    restored = store.undo(ws['id'], ws['revision'])
    raw = store.read(restored['id'])
    restored_folder = store.folder(ws['id']) / raw['generation']
    assert {d['id']: (restored_folder / d['stored_name']).read_bytes() for d in raw['documents']} == original
