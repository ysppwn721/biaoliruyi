"""Bounded OOXML adapters, stable locations, run-preserving patches and real charts."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from docx import Document
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

from .engine import stable_id, number, fmt, convert

HEADERS = ['事实ID', '主体', '指标', '期间', '数值', '单位', '统计口径']


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def validate_office(path):
    if Path(path).suffix.lower() not in ('.xlsx', '.docx', '.pptx'):
        raise ValueError('仅支持 xlsx、docx、pptx 文件')
    try:
        with ZipFile(path) as z:
            infos = z.infolist()
            if len(infos) > 5000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024:
                raise ValueError('文件展开后过大，请拆分为较小文件')
            if any('vbaproject' in i.filename.lower() for i in infos):
                raise ValueError('首版不支持包含宏的文件')
            expected = {'.xlsx': 'xl/workbook.xml', '.docx': 'word/document.xml', '.pptx': 'ppt/presentation.xml'}[Path(path).suffix.lower()]
            if expected not in z.namelist():
                raise ValueError('扩展名与 Office 文件内容不一致')
    except BadZipFile:
        raise ValueError('文件不是有效的 Office 开放 XML 文件')


def read_facts(path, file_id):
    validate_office(path)
    wb = load_workbook(path, data_only=False, read_only=True)
    facts, used = [], set()
    try:
        for sheet in wb.worksheets:
            if sheet.max_row > 2002 or sheet.max_column > 100:
                raise ValueError('首版每张事实表限制为2000行、100列')
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows, []))
            if not all(h in header for h in HEADERS):
                continue
            idx = {h: header.index(h) for h in HEADERS}
            # 只在循环外算一次列字母。原先每行调用 sheet.cell()：在 read_only
            # 模式下每个 cell() 都要重新解析工作表 XML，整体是 O(行数²)——800 行
            # 约 19 秒、1999 行约 2 分钟，用户会以为程序卡死。改用列字母拼坐标后
            # 同等规模耗时降到毫秒级，坐标字符串完全一致。
            cell_column = get_column_letter(idx['数值'] + 1)
            for rno, row in enumerate(rows, 2):
                if not any(v is not None for v in row):
                    continue
                raw = {h: row[i] if i < len(row) else None for h, i in idx.items()}
                fid = str(raw['事实ID'] or '').strip()
                if not fid or fid in used:
                    raise ValueError(f'{sheet.title} 第{rno}行：事实ID为空或重复')
                used.add(fid)
                if len(fid) > 80:
                    raise ValueError('事实ID不能超过80字符')
                for field in ('主体', '指标', '期间', '单位', '统计口径'):
                    if raw[field] is None or not str(raw[field]).strip():
                        raise ValueError(f'{sheet.title} 第{rno}行缺少{field}')
                value = raw['数值']
                if isinstance(value, str) and value.startswith('='):
                    raise ValueError(f'{sheet.title} 第{rno}行使用公式。请先重算并提供事实数值，首版不信任公式缓存')
                if value is not None:
                    value = float(number(value))
                facts.append({'id': fid, 'file_id': file_id, 'subject': str(raw['主体']).strip(),
                              'metric': str(raw['指标']).strip(), 'period': str(raw['期间']).strip(),
                              'value': value, 'unit': str(raw['单位']).strip(), 'scope': str(raw['统计口径']).strip(),
                              'sheet': sheet.title, 'cell': cell_column + str(rno)})
    finally:
        wb.close()
    if not facts:
        raise ValueError('没有找到事实表。首行需包含：' + '、'.join(HEADERS) + '。可下载模板或载入演示项目')
    if len(facts) > 2000:
        raise ValueError('首版单个项目最多2000条事实')
    return facts


def docx_paragraphs(doc):
    for i, para in enumerate(doc.paragraphs):
        yield ['p', i], f'正文第{i+1}段', para
    for ti, table in enumerate(doc.tables):
        seen = set()
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                for pi, para in enumerate(cell.paragraphs):
                    yield ['t', ti, ri, ci, pi], f'表{ti+1} 第{ri+1}行第{ci+1}列', para


def read_document(path, file_id, facts):
    validate_office(path)
    blocks, charts, warnings = [], [], []
    ext = Path(path).suffix.lower()
    if ext == '.docx':
        doc = Document(path)
        entries = docx_paragraphs(doc)
        for loc, label, para in entries:
            if not para.text.strip():
                continue
            if ''.join(r.text for r in para.runs) != para.text:
                warnings.append(label + '含超链接或复杂域，未自动处理')
                continue
            blocks.append({'file_id': file_id, 'location': json.dumps(loc), 'label': label, 'text': para.text})
        if doc.inline_shapes:
            warnings.append('Word内的图片、嵌入图表与图片文字未纳入自动检查')
        if any(s.header.paragraphs[0].text or s.footer.paragraphs[0].text for s in doc.sections):
            warnings.append('页眉页脚未纳入检查')
    elif ext == '.pptx':
        prs = Presentation(path)
        for si, slide in enumerate(prs.slides):
            for sh in slide.shapes:
                if sh.has_text_frame:
                    for pi, para in enumerate(sh.text_frame.paragraphs):
                        if not para.text.strip():
                            continue
                        label = f'第{si+1}页 · {sh.name} · 第{pi+1}段'
                        if ''.join(r.text for r in para.runs) != para.text:
                            warnings.append(label + '包含软换行，未自动处理')
                            continue
                        blocks.append({'file_id': file_id, 'location': json.dumps(['p', si, sh.shape_id, pi]), 'label': label, 'text': para.text})
                if sh.has_table:
                    for ri, row in enumerate(sh.table.rows):
                        for ci, cell in enumerate(row.cells):
                            for pi, para in enumerate(cell.text_frame.paragraphs):
                                if para.text.strip() and ''.join(r.text for r in para.runs) == para.text:
                                    blocks.append({'file_id': file_id, 'location': json.dumps(['t', si, sh.shape_id, ri, ci, pi]),
                                                   'label': f'第{si+1}页表格 第{ri+1}行第{ci+1}列', 'text': para.text})
                if sh.has_chart:
                    chart = sh.chart
                    if chart.chart_type != XL_CHART_TYPE.COLUMN_CLUSTERED or len(chart.series) != 1:
                        warnings.append(f'第{si+1}页图表不属于单系列簇状柱形图，需人工复核')
                        continue
                    categories = [str(c.label) for c in chart.plots[0].categories]
                    series = chart.series[0]
                    refs = []
                    for cat in categories:
                        options = [f for f in facts if f['metric'] in series.name and f['period'] == cat]
                        if len(options) == 1:
                            refs.append(options[0]['id'])
                    unit = next((f['unit'] for f in facts if f['id'] in refs), '')
                    values = list(series.values)
                    if len(refs) != len(values) or any(v is None for v in values):
                        warnings.append(f'第{si+1}页图表无法唯一对应事实及单位，需人工复核')
                        continue
                    loc = json.dumps(['chart', si, sh.shape_id])
                    charts.append({'id': stable_id(file_id, loc, 'chart'), 'file_id': file_id,
                                   'location': loc, 'label': f'第{si+1}页 · 原生柱状图', 'kind': 'chart', 'refs': refs,
                                   'original': series.name + '：' + '，'.join(f'{c} {fmt(v)}{unit}' for c, v in zip(categories, values)),
                                   'spec': {'series': series.name, 'categories': categories, 'values': values, 'unit': unit},
                                   'confirmed': False, 'extraction': '图表结构识别', 'issue': '', 'start': 0, 'end': 0})
                if sh.shape_type in (6, 13):
                    warnings.append(f'第{si+1}页包含组合对象或图片，未检查其内部内容')
        warnings.append('幻灯片备注及母版中的文字未纳入检查')
    else:
        raise ValueError('成果文件只支持Word和PPT')
    if len(blocks) > 2000:
        raise ValueError('成果文件内容超过首版限制（2000个文本块）')
    return blocks, charts, warnings


def patch_runs(para, edits):
    """Replace character spans in reverse order, retaining surrounding run formatting."""
    for start, end, expected, replacement in sorted(edits, reverse=True):
        full = ''.join(r.text for r in para.runs)
        if full[start:end] != expected:
            raise ValueError('原文锚点已变化，已停止写入，请重新分析')
        offset, inserted = 0, False
        for run in para.runs:
            text = run.text
            left, right = offset, offset + len(text)
            offset = right
            if right <= start or left >= end:
                continue
            a, b = max(0, start-left), min(len(text), end-left)
            run.text = text[:a] + (replacement if not inserted else '') + text[b:]
            inserted = True
        if not inserted:
            raise ValueError('无法定位原文字符范围')


def apply_document(source, destination, patches, facts):
    ext = Path(source).suffix.lower()
    document = Document(source) if ext == '.docx' else Presentation(source)
    groups = {}
    for patch in patches:
        groups.setdefault(patch['location'], []).append(patch)
    byid = {f['id']: f for f in facts}
    for location, group in groups.items():
        loc = json.loads(location)
        if ext == '.docx':
            if loc[0] == 'p':
                para = document.paragraphs[loc[1]]
            else:
                para = document.tables[loc[1]].cell(loc[2], loc[3]).paragraphs[loc[4]]
        else:
            slide = document.slides[loc[1]]
            shape = next((s for s in slide.shapes if s.shape_id == loc[2]), None)
            if shape is None:
                raise ValueError('找不到原幻灯片对象')
            if loc[0] == 'chart':
                patch = group[0]
                s = patch['spec']
                chart_data = CategoryChartData()
                chart_data.categories = s['categories']
                chart_data.add_series(s['series'], [float(convert(byid[i]['value'], byid[i]['unit'], s['unit'])) for i in patch['refs']])
                shape.chart.replace_data(chart_data)
                continue
            para = shape.text_frame.paragraphs[loc[3]] if loc[0] == 'p' else shape.table.cell(loc[3], loc[4]).text_frame.paragraphs[loc[5]]
        edits = [(p['start'], p['end'], p['original'], p['expected']) for p in group]
        # Overlapping edits are unsafe even if their replacement happens to look correct.
        ordered = sorted(edits)
        if any(a[1] > b[0] for a, b in zip(ordered, ordered[1:])):
            raise ValueError('检测到重叠修改，需人工复核')
        patch_runs(para, edits)
    document.save(destination)


def update_workbook(source, destination, facts, changes):
    wb = load_workbook(source)
    try:
        for fact in facts:
            if fact['id'] in changes:
                wb[fact['sheet']][fact['cell']] = float(number(changes[fact['id']]))
        wb.save(destination)
    finally:
        wb.close()
