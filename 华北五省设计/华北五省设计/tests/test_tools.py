import json

from openpyxl import load_workbook

from zhilian import tools
from zhilian.demo import create_demo
from zhilian.mcp_server import build_mcp
from zhilian.store import Store


EXPECTED = {'list_projects', 'parse_facts', 'parse_document', 'verify_claim',
            'verify_project', 'diagnose_project', 'diff_excel'}


def test_tool_registry_and_mcp_registration():
    assert len(tools.TOOLS) == 7
    assert {item['name'] for item in tools.TOOLS} == EXPECTED
    server = build_mcp()
    assert {item.name for item in server._tool_manager.list_tools()} == EXPECTED


def test_read_only_tools_and_json_serialization(tmp_path):
    paths = create_demo(tmp_path / 'files')
    store = Store(tmp_path / 'data')
    ws = store.create('工具测试', paths, demo=True)

    listed = tools.invoke(store, 'list_projects', {})
    assert listed['projects'][0]['id'] == ws['id']
    assert {'id', 'name', 'revision'} <= set(listed['projects'][0])

    facts = tools.invoke(store, 'parse_facts', {'excel_path': str(paths[0])})
    parsed = tools.invoke(store, 'parse_document', {
        'excel_path': str(paths[0]), 'document_path': str(paths[1])})
    assert facts['facts']
    assert parsed['claims']
    assert isinstance(parsed['warnings'], list)

    checked = tools.invoke(store, 'verify_claim', {
        'claim': ws['claims'][0], 'facts': ws['facts']})
    assert checked['status'] in {'consistent', 'inconsistent', 'unverifiable'}
    assert {'code', 'evidence'} <= set(checked)

    verified = tools.invoke(store, 'verify_project', {'project_id': ws['id']})
    diagnosed = tools.invoke(store, 'diagnose_project', {'project_id': ws['id']})
    assert verified['summary']['claims'] > 0
    assert isinstance(diagnosed['records'], list)
    assert {'total', 'consistent', 'inconsistent'} <= set(diagnosed['summary'])

    changed = tmp_path / 'changed.xlsx'
    changed.write_bytes(paths[0].read_bytes())
    workbook = load_workbook(changed)
    workbook.active['E3'] = 90
    workbook.save(changed)
    diff = tools.invoke(store, 'diff_excel', {'project_id': ws['id'], 'excel_path': str(changed)})
    assert diff['changes']
    assert diff['summary']['affected_claims'] >= 0
    assert store.read(ws['id'])['revision'] == ws['revision']

    for name in EXPECTED:
        params = {
            'list_projects': {},
            'parse_facts': {'excel_path': str(paths[0])},
            'parse_document': {'excel_path': str(paths[0]), 'document_path': str(paths[1])},
            'verify_claim': {'claim': ws['claims'][0], 'facts': ws['facts']},
            'verify_project': {'project_id': ws['id']},
            'diagnose_project': {'project_id': ws['id']},
            'diff_excel': {'project_id': ws['id'], 'excel_path': str(changed)},
        }[name]
        json.dumps(tools.invoke(store, name, params), ensure_ascii=False)


def test_handlers_return_structured_errors(tmp_path):
    store = Store(tmp_path / 'data')
    assert 'error' in tools.invoke(store, 'parse_facts', {'excel_path': str(tmp_path / 'missing.xlsx')})
    assert 'error' in tools.invoke(store, 'verify_project', {'project_id': '0' * 32})
