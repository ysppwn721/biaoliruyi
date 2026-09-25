from copy import deepcopy
from zipfile import ZipFile

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from zhilian import app as app_module
from zhilian.demo import create_demo
from zhilian.report import build_report
from zhilian.store import Store


@pytest.fixture(autouse=True)
def isolated_environment(monkeypatch):
    monkeypatch.setattr(app_module, 'load_environment', lambda: None)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'test-secret')
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', '')
    # 固定生成时刻以比较接口、导出和直接生成的全文。
    monkeypatch.setattr('zhilian.report.now', lambda: '2026-09-20T06:00:00+00:00')
    def forbidden(*args, **kwargs):
        pytest.fail('生成报告不应调用外部模型')
    monkeypatch.setattr(httpx, 'post', forbidden)


@pytest.fixture
def project(tmp_path):
    store = Store(tmp_path / 'data')
    return store, store.create('成果报告测试', create_demo(tmp_path / 'files'), True)


def confirm_all(store, ws):
    return store.confirm(ws['id'], ws['revision'], [
        {'claim_id': c['id'], 'refs': c['refs']} for c in ws['claims']])


def test_report_contains_changes_comparison_evidence_and_todos(project):
    store, ws = project
    ws = store.change(ws['id'], ws['revision'], {'sales_current': 90})
    claim = next(c for c in ws['claims'] if c['kind'] == 'quote')
    ws = store.confirm(ws['id'], ws['revision'], [{'claim_id': claim['id'], 'refs': claim['refs']}])
    ws = store.repair(ws['id'], ws['revision'], [claim['id']])
    before = deepcopy(ws)
    report = build_report(ws)
    for value in ('# 知链成果与变更报告', '销售额', 'sales_current：125.0 → 90.0',
                  '修改前后对照', '原文：本期销售额为125万元', '改为：本期销售额为90万元',
                  '依据：', '事实表!E3', '演示业务', '人工待办', '未确认来源', '本次修复 1 项'):
        assert value in report
    assert ws == before


def test_empty_states_and_completed_project(project):
    store, ws = project
    report = build_report(ws)
    assert '本次无数据变更' in report and '本次无文件修复' in report
    assert '无待办，可交付' not in report
    ws = confirm_all(store, ws)
    assert '无待办，可交付' in build_report(ws)


def test_confirmed_but_unrepaired_is_not_deliverable(project):
    store, ws = project
    ws = confirm_all(store, ws)
    ws = store.change(ws['id'], ws['revision'], {'sales_current': 90})
    report = build_report(ws)
    assert '待修复结论' in report
    assert '无待办，可交付' not in report


def test_report_endpoint_is_readonly_protected_and_uncached(project, monkeypatch):
    store, ws = project
    client = TestClient(app_module.create_app(store.root))
    endpoint = f'/api/projects/{ws["id"]}/report'
    state = (store.folder(ws['id']) / 'state.json').read_bytes()
    response = client.get(endpoint, headers={'Origin': 'https://example.invalid'})
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    assert response.json() == {'revision': ws['revision'], 'report': build_report(ws)}
    assert (store.folder(ws['id']) / 'state.json').read_bytes() == state
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', 'report-password')
    assert client.get(endpoint).status_code == 401
    assert client.get(endpoint, auth=('zhilian', 'wrong')).status_code == 401
    assert client.get(endpoint, auth=('zhilian', 'report-password')).status_code == 200
    assert client.post(endpoint, auth=('zhilian', 'report-password')).status_code == 405


def test_archive_report_matches_without_changing_file_set(project):
    store, ws = project
    ws = confirm_all(store, ws)
    ws = store.change(ws['id'], ws['revision'], {'sales_current': 90})
    ws = store.repair(ws['id'], ws['revision'], [r['claim_id'] for r in ws['checks'] if r['status'] == 'inconsistent'])
    expected = build_report(ws)
    with ZipFile(store.archive(ws['id'])) as archive:
        assert set(archive.namelist()) == {d['name'] for d in ws['documents']} | {'核验报告.md', '知链核验记录.json'}
        assert archive.read('核验报告.md').decode('utf-8') == expected
    assert build_report(store.read(ws['id'])) == expected


def test_imported_data_history_is_included(project, tmp_path):
    store, ws = project
    raw = store.read(ws['id'])
    source = next(d for d in raw['documents'] if d['kind'] == 'xlsx')
    wb = load_workbook(store.folder(ws['id']) / raw['generation'] / source['stored_name'])
    wb.active['E3'] = 90
    incoming = tmp_path / 'updated.xlsx'
    wb.save(incoming)
    wb.close()
    ws = store.import_source(ws['id'], ws['revision'], incoming)
    report = build_report(ws)
    assert 'sales_current：125.0 → 90.0' in report and '事实表!E3' in report


@pytest.mark.parametrize('case,title', [('missing', '数据缺失'), ('count', '来源冲突'),
                                      ('scope', '口径不一致'), ('zero', '需人工复核')])
def test_unverifiable_todos_are_grouped(project, case, title):
    store, ws = project
    raw = store.read(ws['id'])
    claim = next(c for c in raw['claims'] if c['kind'] == ('quote' if case in ('missing', 'count') else 'growth'))
    if case == 'missing':
        claim['refs'] = []
    elif case == 'count':
        claim['refs'] = ['sales_prev', 'sales_current']
    elif case == 'scope':
        claim['refs'].reverse()
    else:
        next(f for f in raw['facts'] if f['id'] == 'sales_prev')['value'] = 0
    report = build_report(raw)
    assert f'### 无法判断 · {title}' in report
    assert f'- [ ] 无法判断：{claim["original"]}' in report


def test_ocr_todos_are_present_without_double_counting(project):
    store, ws = project
    ws = confirm_all(store, ws)
    raw = store.read(ws['id'])
    image_claim = dict(next(c for c in raw['claims'] if c['kind'] == 'quote'),
                       id='image-claim', source='ocr', image_id='image', confirmed=False, repairable=False)
    raw['ocr_claims'] = [image_claim]
    before = deepcopy(raw)
    report = build_report(raw)
    assert '检查论断 12 项' in report and '来源待确认 1 项' in report
    assert report.count('- [ ] 图片论断（只读）') == 1
    assert '需编辑原图后重新导入' in report
    assert '无待办，可交付' not in report
    assert report == build_report(store.public(raw))
    assert raw == before


def test_external_model_text_is_not_in_report(project):
    _, ws = project
    ws['diagnosis_explanations'] = {'anything': 'test-secret provider-error'}
    ws['audit'].append({'time': 'now', 'event': '模型错误', 'detail': 'test-secret provider-error'})
    ws['agent'] = {'error': 'test-secret provider-error'}
    report = build_report(ws)
    assert 'test-secret' not in report and 'provider-error' not in report
