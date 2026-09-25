"""Reproducible synthetic long-Word fixture and HTTP workflow benchmark.

Run from the repository with the bundled Python runtime. Output text is checked
against an author-time manifest, independently of the extraction/check engine.
"""
from pathlib import Path
import json
import shutil
import sys
import tempfile
import time
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from zhilian.office import docx_paragraphs, read_facts
from zhilian.store import Store

OUT = ROOT / '长文Word测试'
QA = ROOT / '.codex_artifacts' / 'long_word'
TOPICS = ['经营概况', '销售复盘', '产品组合', '预算执行', '订单管理', '渠道协同',
          '客户服务', '交付安排', '库存管理', '回款跟踪', '成本记录', '活动复盘',
          '质量反馈', '供应沟通', '区域协作', '人员安排', '培训记录', '信息归档',
          '需求评审', '风险记录', '计划调整', '会议纪要', '成果交接', '后续跟踪']


def style_font(style, size):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def build():
    OUT.mkdir(exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    # A read-only copy of the original six-fact fixture, not the user's edited file.
    origin = ROOT / '.zhilian/0202ab113e1849bb959142f8757c67f1/v0/d3aaf21691aa4180.xlsx'
    shutil.copy2(origin, OUT / '配套数据_初始.xlsx')
    facts = {f['id']: f['value'] for f in read_facts(OUT / '配套数据_初始.xlsx', 'fixture')}
    assert facts == dict(sales_prev=100, sales_current=125, product_a=80, product_b=65, spending=80, budget=100)
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(.7)
    s.left_margin = s.right_margin = Inches(.8)
    for name, size in [('Normal', 11), ('Title', 21), ('Heading 1', 17), ('Heading 2', 12)]:
        style_font(doc.styles[name], size)
        for b in list(doc.styles[name].element.iter(qn('w:pBdr'))):
            b.getparent().remove(b)
    normal = doc.styles['Normal'].paragraph_format
    # Keep 11pt body text readable while allowing each deliberately page-broken
    # chapter (including its small table) to fit on one page in WPS.
    normal.line_spacing = 1.15
    normal.space_after = Pt(4)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(6)
    s.header.paragraphs[0].text = '知链长文测试  模拟经营报告'
    style_font(doc.styles['Header'], 9)
    s.footer.paragraphs[0].text = '版本编号 125  保留的页脚信息'
    style_font(doc.styles['Footer'], 9)
    planned = []
    def add(text, mild=None, severe=None, count=0, para=None):
        p = para if para is not None else doc.add_paragraph()
        # Split across runs so replacements must preserve surrounding XML styles.
        if '125万元' in text:
            a, b = text.split('125万元', 1)
            p.add_run(a)
            p.add_run('12').bold = True
            p.add_run('5').italic = True
            p.add_run('万元' + b)
        else:
            p.add_run(text)
        planned.append((p._p, text, mild or text, severe or text, count))
        return p

    for n, topic in enumerate(TOPICS, 1):
        h = doc.add_heading('经营报告长文本测试' if n == 1 else f'第{n:02d}章 {topic}', 0 if n == 1 else 1)
        if n > 1:
            h.paragraph_format.page_break_before = True
        if n == 1:
            doc.add_paragraph('模拟数据用于验证跨章节引用、长段落与普通表格的更新。各章使用同一份事实表，不代表不同部门的独立收入。')
        doc.add_heading(f'{topic}记录与核对', 2)
        add(f'第{n:02d}章围绕{topic}整理本期工作的背景、数据引用和后续安排。各项记录采用统一的期间标签，业务说明与数值引用分别维护。阅读本章时，应结合原始记录理解指标范围，不把重复出现的数据累加为新的合计。本文使用模拟资料，用于检查长文更新后的定位是否准确。')
        lead = f'在{topic}讨论中，记录人员先核对来源表的期间和单位，再把已核实的指标写入正文。重复引用服务于不同的阅读场景，数值仍来自相同事实。'
        add(lead + '本期销售额为125万元，较上期增长25%，且超过120万元。',
            lead + '本期销售额为123万元，较上期增长23%，且超过120万元。',
            lead + '本期销售额为90万元，较上期下降10%，且未超过120万元。', 3)
        add(f'{topic}的业务解释需要结合具体过程记录。处理人员应保留需求提出、审批、执行和结果反馈之间的联系，尤其要记录数据尚未覆盖的情况。如果某项观察只有口头描述，应在后续复核时补充出处，不能仅根据相邻的数字推导其原因。')
        add('本期销售额较上期增长。本期销售额超过80万元。',
            severe='本期销售额较上期下降。本期销售额超过80万元。', count=2)
        add('A产品销量最高。支出未超过预算。',
            severe='B产品销量最高。支出超过预算。', count=2)
        add('本期支出为80万元。本期预算为100万元。',
            severe='本期支出为110万元。本期预算为100万元。', count=2)
        add(f'本章的会议资料编号为125，归档批次为2026，附件页码从80开始。这些数字是背景信息，不对应事实表中的经营指标。记录人员负责保留原有措辞，数值调整时只处理已经明确关联的内容，避免把编号、页码或者说明性的数字一并替换。')
        add('本期销售额为125万元；本期支出低于100万元。',
            '本期销售额为123万元；本期支出低于100万元。',
            '本期销售额为90万元；本期支出不少于100万元。', 2)
        if n % 3 == 0:
            table = doc.add_table(rows=1, cols=2)
            table.autofit = False
            table.columns[0].width, table.columns[1].width = Inches(1.2), Inches(5.7)
            for c, t in zip(table.rows[0].cells, ['核对事项', '引用内容']):
                c.text = t
            hdr = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(hdr)
            for label, text, mild, severe in [
                ('销售金额', '本期销售额为125万元。', '本期销售额为123万元。', '本期销售额为90万元。'),
                ('产品比较', 'A产品销量最高。', 'A产品销量最高。', 'B产品销量最高。')]:
                cells = table.add_row().cells
                cells[0].text = label
                add(text, mild, severe, 1, cells[1].paragraphs[0])
            for row_id, row in enumerate(table.rows):
                for c in row.cells:
                    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr = c._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for edge in ['top', 'left', 'bottom', 'right']:
                        e = OxmlElement('w:'+edge)
                        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:'+k),v)
                        borders.append(e)
                    pr.append(borders)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if row_id==0 else 'FFFFFF');pr.append(shade)
                    margins=OxmlElement('w:tcMar')
                    for edge in ['top','bottom','left','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
                    pr.append(margins)
                    for p in c.paragraphs:
                        p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
        add(f'关于{topic}的跟进安排，需要保留现场记录和说明材料。团队会在下一次复核时检查待办事项是否闭环，并把无法从表格计算的观察交由记录人员核实。这里的业务意见属于背景文字，既不应因数值变更被整段重写，也不能因为同页的数值已经通过检查而被视为得到证明。')
        add('人工复核区：本期销售额约为一百二十五万元，增长主要来自渠道优化。此表达保留原文，由人工复核。')
    original = OUT / '长文测试_原始报告.docx'
    doc.save(original)
    locations = {p._p: json.dumps(loc) for loc, _, p in docx_paragraphs(doc)}
    manifest = [{'location':locations[p], 'before':b, 'mild':m, 'severe':s, 'claim_count':c} for p,b,m,s,c in planned]
    (QA/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    return manifest


def benchmark(manifest):
    from fastapi.testclient import TestClient
    from zhilian.app import create_app
    reports=[]
    for scenario, changes in [('mild',{'sales_current':123}), ('severe',{'sales_current':90,'product_a':60,'spending':110})]:
        with tempfile.TemporaryDirectory() as temp:
            app=create_app(Path(temp)/'http')
            client=TestClient(app)
            timings={}
            def request(stage, method, url, **kwargs):
                started=time.perf_counter()
                r=client.request(method,url,**kwargs)
                timings[stage]=round(time.perf_counter()-started,3)
                assert r.status_code==200,(stage,r.status_code,r.text[:600])
                return r
            files=[OUT/'配套数据_初始.xlsx',OUT/'长文测试_原始报告.docx']
            w=request('import','POST','/api/projects',data={'name':'长文测试 '+scenario},
                      files=[('files',(p.name,p.read_bytes())) for p in files]).json()
            initial_summary=w['summary']
            assert initial_summary['claims']==sum(x['claim_count'] for x in manifest), initial_summary
            assert initial_summary['consistent']==initial_summary['claims'],initial_summary
            assert len({c['id'] for c in w['claims']})==len(w['claims'])
            for m in manifest:
                assert sum(c['location']==m['location'] for c in w['claims'])==m['claim_count'],m
            endpoint='/api/projects/'+w['id']
            w=request('confirm','POST',endpoint+'/links',json={'revision':w['revision'],'links':[
                {'claim_id':c['id'],'refs':c['refs']} for c in w['claims']]}).json()
            # Produce the incoming fixture through the existing application, as a user would.
            helper=Store(Path(temp)/'source')
            hw=helper.create('source fixture',files)
            hw=helper.change(hw['id'],hw['revision'],changes)
            hs=helper.read(hw['id']);sd=next(d for d in hs['documents'] if d['kind']=='xlsx')
            incoming=helper.folder(hw['id'])/hs['generation']/sd['stored_name']
            target=OUT/('配套数据_轻微变化.xlsx' if scenario=='mild' else '配套数据_更新后.xlsx')
            shutil.copy2(incoming,target)
            upload={'data':{'revision':w['revision']},'files':{'file':(target.name,target.read_bytes())}}
            preview=request('preview','POST',endpoint+'/source/preview',**upload).json()
            w=request('update','POST',endpoint+'/source',**upload).json()
            before_repair=w['summary']
            ids=[r['claim_id'] for r in w['checks'] if r['status']=='inconsistent']
            w=request('repair','POST',endpoint+'/repair',json={'revision':w['revision'],'claim_ids':ids}).json()
            assert w['summary']['inconsistent']==0 and w['summary']['pending']==0,w['summary']
            d=next(d for d in w['documents'] if d['kind']=='docx')
            content=request('download','GET',d['download_url']).content
            dest=OUT/('长文测试_轻微变化结果.docx' if scenario=='mild' else '长文测试_修复结果.docx')
            dest.write_bytes(content)
            actual={json.dumps(loc):p.text for loc,_,p in docx_paragraphs(Document(dest))}
            expected={json.dumps(loc):p.text for loc,_,p in docx_paragraphs(Document(files[1]))}
            for m in manifest: expected[m['location']]=m[scenario]
            assert actual==expected,'Independent full-document text mismatch'
            # Check table/paragraph structure, styles, sections and unmodified runs.
            old=Document(files[1]);new=Document(dest)
            assert len(old.tables)==len(new.tables)==8
            untouched_runs=0
            for (_,_,p),(_,_,q) in zip(docx_paragraphs(old),docx_paragraphs(new)):
                assert p.style.style_id==q.style.style_id
                assert (p._p.pPr.xml if p._p.pPr is not None else None)==(q._p.pPr.xml if q._p.pPr is not None else None)
                assert len(p.runs)==len(q.runs)
                if p.text==q.text:
                    assert p._p.xml==q._p.xml
                    untouched_runs+=len(p.runs)
                for a,b in zip(p.runs,q.runs):
                    assert (a._r.rPr.xml if a._r.rPr is not None else None)==(b._r.rPr.xml if b._r.rPr is not None else None)
            with ZipFile(files[1]) as a,ZipFile(dest) as b:
                same_parts=[n for n in a.namelist() if n!='word/document.xml']
                assert a.namelist()==b.namelist()
                assert all(a.read(n)==b.read(n) for n in same_parts)
            repaired_summary=w['summary']
            w=request('undo_repair','POST',endpoint+'/undo',json={'revision':w['revision']}).json()
            assert w['summary']['inconsistent']==before_repair['inconsistent']
            w=request('undo_update','POST',endpoint+'/undo',json={'revision':w['revision']}).json()
            assert w['summary']['consistent']==initial_summary['consistent']
            reports.append({'scenario':scenario,'initial':initial_summary,'after_update':before_repair,
                'after_repair':repaired_summary,'preview':preview['summary'],'timings_seconds':timings,
                'full_text_matches_manifest':True,'untouched_run_count':untouched_runs,
                'non_document_parts_unchanged':len(same_parts),'undo_verified':True})
    doc=Document(OUT/'长文测试_原始报告.docx')
    data={'chapters':24,'tables':8,'characters':sum(len(p.text) for _,_,p in docx_paragraphs(doc)),
          'paragraphs_including_table_cells':sum(1 for _ in docx_paragraphs(doc)),
          'expected_claims':sum(x['claim_count'] for x in manifest),'reports':reports,
          'scope':'Synthetic repeated six-fact dataset. Ordinary body paragraphs and first-level tables only. Manual prose and headers/footers stay unverified.'}
    (QA/'results.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(data,ensure_ascii=True))


if __name__=='__main__':
    if '--build-only' in sys.argv:
        build()
        print('Built original DOCX and independent manifest.')
    else:
        benchmark(json.loads((QA/'manifest.json').read_text(encoding='utf-8')))
