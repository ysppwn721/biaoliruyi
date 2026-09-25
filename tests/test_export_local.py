from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from zhilian import app as app_module
from zhilian.demo import create_demo
from zhilian.store import Store


@pytest.fixture(autouse=True)
def isolated_environment(monkeypatch, tmp_path):
    monkeypatch.setattr(app_module, 'load_environment', lambda: None)
    monkeypatch.setenv('ZHILIAN_ACCESS_PASSWORD', '')
    monkeypatch.setenv('ZHILIAN_OUTPUT_DIR', str(tmp_path / 'outputs'))
    monkeypatch.delenv('ZHILIAN_AUTO_EXPORT', raising=False)


@pytest.fixture
def project(tmp_path):
    store = Store(tmp_path / 'data')
    return store, store.create('本地导出测试', create_demo(tmp_path / 'files'), True)


def finish_repair(store, ws):
    ws = store.confirm(ws['id'], ws['revision'], [
        {'claim_id': c['id'], 'refs': c['refs']} for c in ws['claims']])
    ws = store.change(ws['id'], ws['revision'], {
        'sales_current': 90, 'product_a': 60, 'spending': 110})
    ids = [r['claim_id'] for r in ws['checks'] if r['status'] == 'inconsistent']
    return store.repair(ws['id'], ws['revision'], ids)


def test_local_export_copies_files_and_report(project):
    store, ws = project
    ws = finish_repair(store, ws)
    result = store.export_local(ws['id'])
    folder = Path(result['folder'])
    assert folder.is_dir() and {f['name'] for f in result['files']} == {
        '业务数据.xlsx', '分析报告.docx', '业务汇报.pptx', '变更报告.md'}
    doc = Document(folder / '分析报告.docx')
    assert '本期销售额为90万元' in '\n'.join(p.text for p in doc.paragraphs)
    prs = Presentation(folder / '业务汇报.pptx')
    text = '\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert '较上期下降10%' in text
    wb = load_workbook(folder / '业务数据.xlsx', data_only=True)
    assert wb.active['E3'].value == 90
    wb.close()
    assert '修改前后对照' in (folder / '变更报告.md').read_text(encoding='utf-8')
    raw = store.read(ws['id'])
    assert raw['revision'] == ws['revision'] + 1
    assert raw['audit'][-1]['event'] == '导出到本地'


def test_local_export_revision_isolation_and_same_revision_overwrite(project):
    store, ws = project
    first = store.export_local(ws['id'])
    current = store.read(ws['id'])
    second = store.export_local(ws['id'])
    assert first['folder'] != second['folder']
    # 重置版本仅用于模拟同版本重试，目标目录应安全覆盖。
    raw = store.read(ws['id'])
    raw['revision'] -= 1
    store.write(raw)
    third = store.export_local(ws['id'])
    assert third['folder'] == second['folder']
    assert Path(third['folder']).is_dir()


@pytest.mark.parametrize('name', ['项目<>:"|?*\\/名称', '   '])
def test_local_export_cleans_project_name(project, name):
    store, ws = project
    raw = store.read(ws['id'])
    raw['name'] = name
    store.write(raw)
    result = store.export_local(ws['id'])
    folder = Path(result['folder'])
    assert folder.is_dir() and folder.parent.name == 'outputs'
    assert all(c not in folder.name for c in '<>:"|?*\\/')


def test_local_export_api_and_origin_protection(project):
    store, ws = project
    client = TestClient(app_module.create_app(store.root))
    endpoint = f'/api/projects/{ws["id"]}/export-local'
    response = client.post(endpoint, json={'revision': ws['revision']}, headers={'Origin': 'https://evil.example'})
    assert response.status_code == 403
    response = client.post(endpoint, json={'revision': ws['revision']})
    assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
    body = response.json()
    assert Path(body['folder']).is_dir() and body['files']


def test_local_export_failure_does_not_change_workspace(project, monkeypatch):
    store, ws = project
    before = store.read(ws['id'])
    monkeypatch.setattr('zhilian.store.shutil.copyfile', lambda *args: (_ for _ in ()).throw(OSError('permission denied')))
    with pytest.raises(ValueError, match='无法写入本地输出目录'):
        store.export_local(ws['id'])
    after = store.read(ws['id'])
    assert after == before


def test_local_export_does_not_modify_generation_files(project):
    store, ws = project
    raw = store.read(ws['id'])
    source_root = store.folder(ws['id']) / raw['generation']
    before = {d['id']: (source_root / d['stored_name']).read_bytes() for d in raw['documents']}
    store.export_local(ws['id'])
    raw = store.read(ws['id'])
    source_root = store.folder(ws['id']) / raw['generation']
    assert {d['id']: (source_root / d['stored_name']).read_bytes() for d in raw['documents']} == before
