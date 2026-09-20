from docx import Document

from zhilian.demo import create_demo
from zhilian.engine import check, extract_claims, stable_id, unmatched_spans
from zhilian.store import Store


def test_compound_sentence_extracts_independent_claims(tmp_path):
    files = create_demo(tmp_path / "files")
    store = Store(tmp_path / "data")
    workspace = store.create("compound", files[:2])
    facts = workspace["facts"]
    text = "本期销售额为125万元，较上期增长25%，且超过120万元。"
    claims = extract_claims({"file_id": "f", "location": "p", "label": "段落", "text": text}, facts)

    assert [c["kind"] for c in claims] == ["quote", "growth", "threshold"]
    # The pre-compound extractor treated this sentence as a growth claim. Its
    # stable ID remains intact so an existing confirmation can be migrated.
    assert claims[1]["id"] == stable_id("f", "p", "growth", 0)
    assert all(text[c["start"] : c["end"]] == c["original"] for c in claims)
    assert all(check(c, facts)["status"] == "consistent" for c in claims)
    assert unmatched_spans([{"file_id": "f", "location": "p", "label": "段落", "text": text}], claims) == []

    two_quotes = extract_claims({"file_id": "f", "location": "q", "label": "段落", "text": "销售额为125万元，支出为80万元。"}, facts)
    assert [c["kind"] for c in two_quotes] == ["quote", "quote"]
    assert len({c["id"] for c in two_quotes}) == 2

    two_ranks = extract_claims({"file_id": "f", "location": "r", "label": "段落",
                                "text": "A产品销量最高，B产品销量最高。"}, facts)
    assert [c["original"] for c in two_ranks] == ["A产品销量最高", "B产品销量最高"]


def test_compound_word_claims_can_be_repaired_together(tmp_path):
    paths = create_demo(tmp_path / "files")
    doc = Document(paths[1])
    for para in doc.paragraphs:
        if "A产品销量最高" in para.text:
            para.text = "A产品销量最高，支出未超过预算。"
    doc.save(paths[1])

    store = Store(tmp_path / "data")
    workspace = store.create("compound", paths[:2])
    compound = [c for c in workspace["claims"] if c["original"] in ("A产品销量最高", "支出未超过预算")]
    assert len(compound) == 2
    workspace = store.confirm(workspace["id"], workspace["revision"],
                              [{"claim_id": c["id"], "refs": c["refs"]} for c in workspace["claims"]])
    workspace = store.change(workspace["id"], workspace["revision"], {"product_a": 60, "spending": 110})
    invalid = [c["id"] for c in workspace["claims"]
               if c["original"] in ("A产品销量最高", "支出未超过预算")
               and check(c, workspace["facts"])["status"] == "inconsistent"]
    assert len(invalid) == 2
    workspace = store.repair(workspace["id"], workspace["revision"], invalid)
    assert workspace["summary"]["inconsistent"] == 0

    state = store.read(workspace["id"])
    docx = next(d for d in state["documents"] if d["kind"] == "docx")
    repaired = Document(store.folder(workspace["id"]) / state["generation"] / docx["stored_name"])
    assert any("B产品销量最高，支出超过预算" in para.text for para in repaired.paragraphs)


def test_compound_single_repair_leaves_second_claim_unchanged(tmp_path):
    paths = create_demo(tmp_path / "files")
    doc = Document(paths[1])
    for para in doc.paragraphs:
        if "本期销售额为" in para.text:
            para.text = "本期销售额为125万元，较上期增长25%。"
            break
    doc.save(paths[1])
    store = Store(tmp_path / "data")
    workspace = store.create("compound", paths[:2])
    workspace = store.confirm(workspace["id"], workspace["revision"],
                              [{"claim_id": c["id"], "refs": c["refs"]} for c in workspace["claims"]])
    workspace = store.change(workspace["id"], workspace["revision"], {"sales_current": 90})
    quote = next(c for c in workspace["claims"] if c["kind"] == "quote" and c["file_id"] != next(d for d in workspace["documents"] if d["kind"] == "xlsx")["id"])
    assert check(quote, workspace["facts"])["status"] == "inconsistent"
    before_growth = next(c for c in workspace["claims"] if c["kind"] == "growth" and c["file_id"] == quote["file_id"])
    assert check(before_growth, workspace["facts"])["status"] == "inconsistent"
    workspace = store.repair(workspace["id"], workspace["revision"], [quote["id"]])
    after_growth = next(c for c in workspace["claims"] if c["id"] == before_growth["id"])
    assert check(after_growth, workspace["facts"])["status"] == "inconsistent"
    assert any("本期销售额为90万元，较上期增长25%" in b["text"] for b in workspace["blocks"])
