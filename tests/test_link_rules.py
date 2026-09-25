from copy import deepcopy

import httpx
import pytest
from openpyxl import load_workbook

from zhilian import agent
from zhilian.demo import create_demo
from zhilian.store import Store


KEY = 'quote|总计|销售额|本期'


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setenv('DEEPSEEK_API_KEY', '')
    def unexpected(*args, **kwargs):
        pytest.fail('测试不应发起真实模型请求')
    monkeypatch.setattr(httpx, 'post', unexpected)


@pytest.fixture
def project(tmp_path):
    paths = create_demo(tmp_path / 'files')
    wb = load_workbook(paths[0])
    wb.active.append(['other_sales', '总计', '销售额', '本期', 125, '万元', '其他地区'])
    wb.save(paths[0])
    wb.close()
    store = Store(tmp_path / 'data')
    return store, store.create('关联规则测试', paths)


def quote(ws):
    return next(c for c in ws['claims'] if c['kind'] == 'quote')


def confirm(store, ws, fact_id='sales_current'):
    return store.confirm(ws['id'], ws['revision'], [
        {'claim_id': quote(ws)['id'], 'refs': [fact_id]}])


def rule(ws):
    return next(r for r in ws['link_rules'] if r['key'] == KEY)


def test_agent_confirmation_records_rule_and_audit(project):
    store, ws = project
    assert ws['link_rules'] == []
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    assert ws['agent']['pending']['kind'] == 'resolve_ambiguity'
    cid = quote(ws)['id']
    ws = agent.decide(store, ws['id'], ws['revision'], [{
        'kind': 'resolve_ambiguity', 'items': [{'claim_id': cid, 'refs': ['sales_current']}]}])
    remembered = rule(ws)
    assert remembered['fact_id'] == 'sales_current'
    assert remembered['hits'] == 1
    assert remembered['source_claim_id'] == cid
    assert remembered['scope'] == '演示业务'
    assert remembered['created_at']
    assert any(a['event'] == '记录关联规则' and KEY in a['detail']
               and 'sales_current' in a['detail'] for a in ws['audit'])


def test_options_prioritize_memory_without_changing_claim(project):
    _, ws = project
    claim = quote(ws)
    before = deepcopy(claim)
    options = agent._options(claim, ws['facts'], [], [], [{'key': KEY, 'fact_id': 'other_sales'}])
    assert options[0]['refs'] == ['other_sales']
    assert options[0]['remembered'] is True
    assert options[0]['reason'].startswith('复用已确认规则')
    assert all(o['remembered'] is False for o in options[1:])
    assert claim == before and claim['confirmed'] is False


def test_latest_confirmation_wins(project):
    store, ws = project
    ws = confirm(store, ws)
    old = deepcopy(rule(ws))
    second = next(c for c in ws['claims'] if c['kind'] == 'quote' and c['id'] != quote(ws)['id'])
    ws = store.confirm(ws['id'], ws['revision'], [{'claim_id': second['id'], 'refs': ['other_sales']}])
    updated = rule(ws)
    assert sum(r['key'] == KEY for r in ws['link_rules']) == 1
    assert updated['fact_id'] == 'other_sales' and updated['scope'] == '其他地区'
    assert updated['hits'] == old['hits'] + 1
    assert updated['source_claim_id'] == second['id']
    assert updated['created_at'] >= old['created_at']


def test_rules_survive_restart_and_data_change(project):
    store, ws = project
    ws = confirm(store, ws)
    saved = deepcopy(ws['link_rules'])
    restarted = Store(store.root)
    assert restarted.read(ws['id'])['link_rules'] == saved
    ws = restarted.change(ws['id'], ws['revision'], {'sales_current': 120})
    assert ws['link_rules'] == saved
    assert quote(ws)['confirmed'] and quote(ws)['refs'] == ['sales_current']


@pytest.mark.parametrize('rules', [
    [{'key': KEY, 'fact_id': 'missing'}], [{'key': KEY}], [{}], [None], None,
    {'invalid': 'container'}, [{'key': [], 'fact_id': 'sales_current'}],
])
def test_invalid_rules_are_ignored(project, rules):
    _, ws = project
    options = agent._options(quote(ws), ws['facts'], [], [], rules)
    assert len(options) >= 2
    assert not any(o['remembered'] for o in options)


def test_memory_does_not_skip_ambiguity_gate(project):
    store, ws = project
    ws = confirm(store, ws, 'other_sales')
    saved = deepcopy(ws['link_rules'])
    ws = agent.run_agent(store, ws['id'], ws['revision'])
    pending = ws['agent']['pending']
    assert pending['kind'] == 'resolve_ambiguity'
    item = next(i for i in pending['items'] if i['kind'] == 'quote')
    assert item['claim_id'] != quote(ws)['id']
    assert len(item['options']) >= 2
    assert item['options'][0]['refs'] == ['other_sales']
    assert item['options'][0]['remembered'] is True
    assert not next(c for c in ws['claims'] if c['id'] == item['claim_id'])['confirmed']
    assert ws['link_rules'] == saved


def test_rejected_confirmation_leaves_rules_unchanged(project):
    store, ws = project
    before = store.read(ws['id'])
    with pytest.raises(ValueError, match='无法确认'):
        store.confirm(ws['id'], ws['revision'], [
            {'claim_id': quote(ws)['id'], 'refs': ['sales_current', 'other_sales']}])
    assert store.read(ws['id']) == before


def test_legacy_project_without_rules_can_confirm(project):
    store, ws = project
    raw = store.read(ws['id'])
    raw.pop('link_rules')
    store.write(raw)
    ws = confirm(store, ws)
    assert rule(ws)['hits'] == 1


@pytest.mark.parametrize('field', ['id', 'subject', 'metric', 'period', 'scope'])
def test_missing_metadata_does_not_record_rule(project, field):
    store, ws = project
    fact = deepcopy(next(f for f in ws['facts'] if f['id'] == 'sales_current'))
    fact.pop(field)
    before = deepcopy(ws)
    assert store._remember(ws, quote(ws), fact) is None
    assert ws == before


def test_duplicate_keys_and_invalid_hits_do_not_block_confirmation(project):
    store, ws = project
    raw = store.read(ws['id'])
    raw['link_rules'] = [{'key': KEY, 'fact_id': 'missing', 'hits': 5},
                         {'key': KEY, 'hits': 'invalid'}, None]
    store.write(raw)
    ws = confirm(store, ws, 'other_sales')
    matching = [r for r in ws['link_rules'] if isinstance(r, dict) and r.get('key') == KEY]
    assert len(matching) == 1
    assert matching[0]['fact_id'] == 'other_sales' and matching[0]['hits'] == 1
    assert quote(ws)['confirmed']


def test_multi_fact_option_requires_every_rule(project):
    _, ws = project
    claim = next(c for c in ws['claims'] if c['kind'] == 'growth')
    refs = ['sales_prev', 'sales_current']
    rules = [{'key': 'growth|总计|销售额|上期', 'fact_id': 'sales_prev'}]
    assert not agent._options(claim, ws['facts'], refs, [], rules)[0]['remembered']
    rules.append({'key': 'growth|总计|销售额|本期', 'fact_id': 'sales_current'})
    assert agent._options(claim, ws['facts'], refs, [], rules)[0]['remembered']
